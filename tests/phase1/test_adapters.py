import asyncio

import pytest
from docx import Document

from app.orchestration.artifacts import ArtifactRuntime
from app.orchestration.events import EventBus, Trace
from app.orchestration.mcp_runtime import MCPRuntime
from app.orchestration.sandbox import SandboxRequest, SandboxResult, SandboxRuntime, DockerAdapter
from app.orchestration.runtimes import RuntimeRouter
from app.orchestration.models import PlanStep
from app.security.consent import CONSENT_REQUESTS, approve_consent, deny_consent
from app.task.service import task_service


@pytest.mark.parametrize("exit_code,timed_out", [(0, False), (1, False), (124, True)])
def test_sandbox_contract_and_events(exit_code, timed_out):
    requests = []
    class Adapter:
        async def execute(self, request):
            requests.append(request)
            return SandboxResult(stdout="private stdout", stderr="private stderr", exit_code=exit_code, timed_out=timed_out)
    bus = EventBus()
    operation = SandboxRuntime(Adapter(), mode="docker").execute({"code": "print(1)", "timeout": 2}, Trace(bus, "t"))
    if exit_code:
        with pytest.raises(RuntimeError):
            asyncio.run(operation)
    else:
        assert asyncio.run(operation)["stdout"] == "private stdout"
    assert requests[0].timeout == 2
    assert not any("private stdout" in e.model_dump_json() for e in bus.replay("t"))
    assert bus.replay("t")[-1].metadata["exit_code"] == exit_code
    assert bus.replay("t")[-1].metadata["timed_out"] == timed_out


def test_sandbox_mock_never_calls_adapter():
    class Adapter:
        async def execute(self, request):
            raise AssertionError("must not execute")
    bus = EventBus()
    result = asyncio.run(SandboxRuntime(Adapter(), mode="mock").execute({"code": "arbitrary()"}, Trace(bus, "t")))
    assert result["is_mock"] and "not run" in result["stdout"]
    assert all(e.is_mock for e in bus.replay("t"))


def test_docker_timeout_cleanup_and_safety_flags(monkeypatch):
    calls = []
    class Writer:
        def write(self, value):
            assert value == b"print(1)"
        async def drain(self):
            pass
        def close(self):
            pass
    class Process:
        def __init__(self, hangs=False):
            self.returncode = None if hangs else 0
            self.stdin = Writer()
            self.stdout = asyncio.StreamReader()
            self.stderr = asyncio.StreamReader()
            self.stdout.feed_eof()
            self.stderr.feed_eof()
        async def wait(self):
            if self.returncode is None:
                await asyncio.Event().wait()
            return self.returncode
        def kill(self):
            self.returncode = -9
    async def spawn(*args, **kwargs):
        calls.append(args)
        return Process(hangs=args[1] == "run")
    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    result = asyncio.run(DockerAdapter().execute(SandboxRequest(code="print(1)", timeout=.01)))
    assert result.timed_out and result.exit_code == 124
    for flag in ["--network=none", "--read-only", "--pull=never", "--cap-drop=ALL", "--pids-limit=32", "--user=65534:65534"]:
        assert flag in calls[0]
    assert calls[-1][:3] == ("docker", "rm", "-f")


def test_real_local_mcp_stdio_and_failure_mapping():
    async def scenario():
        runtime = MCPRuntime()
        bus = EventBus()
        trace = Trace(bus, "t", "mcp", 1)
        result = await runtime.call("calculator", {"a": 6, "b": 7, "operation": "multiply"}, trace)
        assert result["value"] == 42
        search = await runtime.call("search_internal_docs", {"query": "inspection"}, trace)
        assert search["result_count"] >= 1
        with pytest.raises(RuntimeError, match="^MCP tool failed$"):
            await runtime.call("calculator", {"a": 1, "b": 2, "operation": "invalid"}, trace)
        assert [e.event_type for e in bus.replay("t")].count("mcp.completed") == 2
        assert bus.replay("t")[-1].event_type == "mcp.failed"
        assert not any(e.is_mock for e in bus.replay("t"))
    asyncio.run(scenario())


def test_real_docx_structure_and_containment(workspace):
    bus = EventBus()
    artifact = ArtifactRuntime().create_word("task_safe", {
        "title": "Review", "content": "Finding", "headings": ["Recommendations"],
        "bullets": ["Replace seal"], "table": [["Owner", "Action"], ["Engineer", "Review"]],
    }, Trace(bus, "task_safe"))
    path = workspace / artifact["path"]
    assert path.is_relative_to(workspace / "outputs" / "tasks" / "task_safe")
    doc = Document(path)
    assert doc.tables[0].cell(1, 1).text == "Review"
    assert "Replace seal" in [p.text for p in doc.paragraphs]
    assert bus.replay("task_safe")[-1].event_type == "artifact.created"
    with pytest.raises(PermissionError):
        ArtifactRuntime().create_word("../../escape", {}, Trace(bus, "bad"))


@pytest.mark.parametrize("approved", [True, False])
def test_consent_wait_resume_or_deny(user, approved):
    async def scenario():
        task = task_service.create_task(user.id, "sandbox")
        step = PlanStep(id="run", title="Run", kind="sandbox", tool_name="python")
        bus = EventBus()
        router = RuntimeRouter(sandbox=SandboxRuntime(mode="mock"))
        work = asyncio.create_task(router.execute(step, user, task.id, {"code": "print(1)"}, Trace(bus, task.id), True))
        await asyncio.sleep(0)
        assert step.status == "waiting"
        consent = next(c for c in CONSENT_REQUESTS.values() if c.task_id == task.id)
        (approve_consent if approved else deny_consent)(consent.id, user)
        if approved:
            assert (await work)["is_mock"]
        else:
            with pytest.raises(PermissionError):
                await work
    asyncio.run(scenario())


@pytest.mark.docker
def test_optional_real_docker():
    import os
    if os.environ.get("SPEED_TEST_DOCKER") != "1":
        pytest.skip("Explicit optional Docker integration only")
    result = asyncio.run(DockerAdapter().execute(SandboxRequest(code="print(6 * 7)")))
    assert result.exit_code == 0 and result.stdout.strip() == "42"


def test_mcp_client_cleanup_on_error(monkeypatch):
    import mcp
    lifecycle = []
    class Client:
        def __init__(self, parameters):
            assert parameters.command
            assert parameters.args[0].endswith("demo_mcp.py")
        async def __aenter__(self):
            lifecycle.append("entered")
            return self
        async def __aexit__(self, *args):
            lifecycle.append("closed")
        async def list_tools(self):
            from types import SimpleNamespace
            return SimpleNamespace(tools=[SimpleNamespace(name="calculator")])
        async def call_tool(self, *args):
            raise RuntimeError("private server failure")
    monkeypatch.setattr(mcp, "Client", Client)
    bus = EventBus()
    with pytest.raises(RuntimeError):
        asyncio.run(MCPRuntime().call("calculator", {"a": 1, "b": 2}, Trace(bus, "t")))
    assert lifecycle == ["entered", "closed"]
    assert bus.replay("t")[-1].event_type == "mcp.failed"
    assert "private server failure" not in str(bus.replay("t"))


@pytest.mark.parametrize("missing_tool,is_error,content,cleanup_fails", [
    (False, True, None, False),
    (False, False, None, False),
    (True, False, None, False),
    (False, True, None, True),
])
def test_mcp_mapped_failure_after_clean_exit(monkeypatch, missing_tool, is_error, content, cleanup_fails):
    import mcp
    from types import SimpleNamespace
    exits = []

    class Client:
        def __init__(self, parameters):
            pass
        async def __aenter__(self):
            return self
        async def __aexit__(self, exc_type, exc, tb):
            exits.append(exc_type)
            if cleanup_fails:
                raise OSError("cleanup failed")
        async def list_tools(self):
            return SimpleNamespace(tools=[] if missing_tool else [SimpleNamespace(name="calculator")])
        async def call_tool(self, *args):
            return SimpleNamespace(is_error=is_error, structured_content=content)

    monkeypatch.setattr(mcp, "Client", Client)
    expected = OSError if cleanup_fails else ValueError if missing_tool else RuntimeError
    bus = EventBus()
    with pytest.raises(expected):
        asyncio.run(MCPRuntime().call("calculator", {}, Trace(bus, "t")))
    assert exits == [None]
    assert [event.event_type for event in bus.replay("t")] == ["mcp.started", "mcp.failed"]
