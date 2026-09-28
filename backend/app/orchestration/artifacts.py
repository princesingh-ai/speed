from uuid import uuid4

from pydantic import BaseModel, Field
from app.tools.builtin.filesystem_scope import filesystem_scope


class WordRequest(BaseModel):
    title: str = Field(default="Approval note", max_length=200)
    content: str = Field(default="", max_length=100_000)
    headings: list[str] = Field(default_factory=list, max_length=30)
    bullets: list[str] = Field(default_factory=list, max_length=100)
    table: list[list[str]] = Field(default_factory=list, max_length=100)


class ArtifactRuntime:
    def create_word(self, task_id, inputs, trace):
        trace.emit("artifact.started", "Generate Word document", "running")
        try:
            from docx import Document
            request = WordRequest.model_validate(inputs)
            artifact_id = uuid4().hex
            relative = f"outputs/tasks/{task_id}/{artifact_id}/approval-note.docx"
            path = filesystem_scope.resolve_workspace_path(relative)
            path.parent.mkdir(parents=True, exist_ok=True)
            document = Document()
            document.add_heading(request.title, 0)
            if trace.is_mock:
                document.add_paragraph("DEMONSTRATION: contains deterministic/mock source material. Human review required.")
            for paragraph in request.content.splitlines():
                document.add_paragraph(paragraph)
            for heading in request.headings:
                document.add_heading(heading, level=1)
            for bullet in request.bullets:
                document.add_paragraph(bullet, style="List Bullet")
            if request.table:
                columns = len(request.table[0])
                if not 1 <= columns <= 12 or any(len(row) != columns for row in request.table):
                    raise ValueError("Table must be rectangular, with 1–12 columns")
                table = document.add_table(rows=0, cols=columns)
                for values in request.table:
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = value
            path = filesystem_scope.resolve_workspace_path(relative)
            document.save(path)
            artifact = {"id": artifact_id, "name": path.name, "path": relative,
                        "size": path.stat().st_size, "is_mock": trace.is_mock}
            trace.emit("artifact.created", "Created approval-note.docx", metadata={
                "artifact_id": artifact_id, "path": relative, "size": artifact["size"]})
            return artifact
        except Exception:
            trace.emit("artifact.failed", "Word generation failed", "failed")
            raise
